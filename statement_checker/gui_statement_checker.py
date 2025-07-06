# statement_checker/gui_statement_checker.py

import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox
from tkcalendar import DateEntry
from datetime import datetime
from collections import defaultdict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from cc_analysis.utils import log_error, log_info  # Use existing logger from main GUI
from cc_analysis.constants import STMT_CHECK_PATH  # Standard output location

def extract_date_from_filename(filename):
    match = re.search(r'(\d{2}-\d{2}-\d{4})', filename)
    if match:
        try:
            return datetime.strptime(match.group(1), "%d-%m-%Y")
        except ValueError:
            return None
    match = re.search(r'_(\d{8})(\d{8})\.pdf$', filename)
    if match:
        try:
            return datetime.strptime(match.group(2), "%d%m%Y")
        except ValueError:
            return None
    return None

def identify_bank(filename):
    if filename.startswith("CC_STMT"):
        return "IndusInd"
    elif filename.startswith("4639") or filename.startswith("4854"):
        return "HDFC"
    return None

def get_month_year_list(start_date, end_date):
    months = []
    current = start_date.replace(day=1)
    while current <= end_date:
        months.append(current.strftime("%b-%Y"))
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1)
        else:
            current = current.replace(month=current.month + 1)
    return months[::-1]

def get_last_date_of_month(month_year_str, bank):
    dt = datetime.strptime(month_year_str, "%b-%Y")
    day = 16 if bank == "HDFC" else 18
    return datetime(dt.year, dt.month, day).strftime("%d-%m-%Y")

def analyze_statements(folder, start_date, end_date):
    data = defaultdict(set)
    debug_log = []
    for filename in os.listdir(folder):
        if not filename.lower().endswith(".pdf"):
            continue
        bank = identify_bank(filename)
        if not bank:
            continue
        date_obj = extract_date_from_filename(filename)
        if not date_obj:
            debug_log.append((filename, "date not found"))
            continue
        if not (start_date <= date_obj <= end_date):
            continue
        month = date_obj.strftime("%b-%Y")
        data[bank].add(month)
    all_months = get_month_year_list(start_date, end_date)
    banks = ["HDFC", "IndusInd"]
    availability = []
    for i, month in enumerate(all_months, 1):
        row = {"Sl.No.": i, "Date": month}
        for bank in banks:
            row[bank] = "Available" if month in data[bank] else "Not Available"
        availability.append(row)
    missing = {bank: [] for bank in banks}
    for bank in banks:
        missing_months = sorted(set(all_months) - data[bank], key=lambda x: datetime.strptime(x, "%b-%Y"))
        for i, month in enumerate(missing_months, 1):
            missing[bank].append({"Sl.No.": i, "Missing Month": month, "Last Expected Date (DD-MM-YYYY)": get_last_date_of_month(month, bank)})
    return availability, missing, debug_log

def export_to_excel(folder, availability, missing):
    os.makedirs(os.path.dirname(STMT_CHECK_PATH), exist_ok=True)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Statement Availability"
    df = pd.DataFrame(availability)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws1.append(r)
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    for row in ws1.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            if cell.value == "Available":
                cell.fill = green
            elif cell.value == "Not Available":
                cell.fill = red
            cell.alignment = Alignment(horizontal="center")
    ws2 = wb.create_sheet(title="Missing Statements Detail")
    ws2.append(["Table 1: HDFC Missing Statements"])
    ws2.append([])
    for r in dataframe_to_rows(pd.DataFrame(missing["HDFC"]), index=False, header=True):
        ws2.append(r)
    ws2.append([])
    ws2.append(["Table 2: IndusInd Missing Statements"])
    ws2.append([])
    for r in dataframe_to_rows(pd.DataFrame(missing["IndusInd"]), index=False, header=True):
        ws2.append(r)
    wb.save(STMT_CHECK_PATH)
    return STMT_CHECK_PATH

def export_debug_log(debug_log):
    for fname, reason in debug_log:
        log_error(f"{fname} - {reason}")

def launch_gui():
    root = tk.Tk()
    root.title("Credit Card Statement Analyzer")
    root.geometry("650x400")
    folder_var = tk.StringVar()

    def browse_folder():
        folder = filedialog.askdirectory()
        if folder:
            folder_var.set(folder)
            preview.delete(1.0, tk.END)
            for fname in sorted(os.listdir(folder)):
                if fname.lower().endswith(".pdf") and identify_bank(fname):
                    preview.insert(tk.END, f"{fname}\n")

    def run_analysis():
        folder = folder_var.get()
        if not folder:
            messagebox.showerror("Error", "Please select a folder")
            return
        try:
            start = datetime.strptime(start_date.get(), "%d-%m-%Y")
            end = datetime.strptime(end_date.get(), "%d-%m-%Y")
            if start > end:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Invalid date range")
            return
        availability, missing, debug_log = analyze_statements(folder, start, end)
        if not availability:
            messagebox.showinfo("No Data", "No valid statements found.")
            return
        path = export_to_excel(folder, availability, missing)
        export_debug_log(debug_log)
        msg = f"Statement summary saved to:\n{path}"
        messagebox.showinfo("Exported", msg)

    tk.Label(root, text="Select Folder:").pack(pady=5)
    folder_frame = tk.Frame(root)
    folder_frame.pack()
    tk.Entry(folder_frame, textvariable=folder_var, width=50).pack(side=tk.LEFT, padx=5)
    tk.Button(folder_frame, text="Browse", command=browse_folder).pack(side=tk.LEFT)
    tk.Label(root, text="Select Date Range:").pack(pady=10)
    date_frame = tk.Frame(root)
    date_frame.pack()
    tk.Label(date_frame, text="Start Date:").grid(row=0, column=0, padx=5)
    start_date = DateEntry(date_frame, date_pattern="dd-mm-yyyy")
    start_date.grid(row=0, column=1, padx=5)
    tk.Label(date_frame, text="End Date:").grid(row=0, column=2, padx=5)
    end_date = DateEntry(date_frame, date_pattern="dd-mm-yyyy")
    end_date.grid(row=0, column=3, padx=5)
    tk.Button(root, text="Analyze and Export", command=run_analysis, bg="green", fg="white").pack(pady=10)
    preview_label = tk.Label(root, text="Matching Files in Folder:")
    preview_label.pack()
    preview = tk.Text(root, height=10, width=80)
    preview.pack()
    root.mainloop()