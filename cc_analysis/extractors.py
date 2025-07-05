import re
import sqlite3

import pdfplumber

from cc_analysis.categorizer import assign_category, load_category_mappings
from cc_analysis.constants import CONSOLIDATED_DB_PATH


def extract_hdfc(pdf_path, password):
    # Improved pattern: handles optional whitespace and "Cr"/"CR"
    pattern = re.compile(r"(\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2}:\d{2})\s+(.+?)\s+([\d,]+\.\d{2})\s*(Cr|CR)?$")
    transactions = []
    categories = load_category_mappings()

    with pdfplumber.open(pdf_path, password=password) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            for line in text.split("\n"):
                match = pattern.match(line.strip())
                if match:
                    datetime_str, description, amount_str, credit_flag = match.groups()
                    date = pd.to_datetime(datetime_str, format="%d/%m/%Y %H:%M:%S", errors='coerce').date()
                    amount = float(amount_str.replace(",", ""))
                    debit = 0
                    credit = 0
                    if credit_flag:
                        credit = amount
                    else:
                        debit = amount
                    category = assign_category(description, "", categories)

                    transactions.append([
                        date,                      # Date
                        description.strip(),       # Description
                        "",                        # Merchant
                        category,                  # Category
                        0,                         # Reward Points
                        "HDFC",                    # Bank
                        "CreditCard",              # SourceType
                        debit,
                        credit,
                        None                       # Balance
                    ])

    return pd.DataFrame(transactions, columns=[
        'Date', 'Description', 'Merchant', 'Category',
        'Reward Points', 'Bank', 'SourceType', 'Debit', 'Credit', 'Balance'
    ])


def extract_indusind(pdf_path, password):
    pattern = re.compile(
        r"(\d{2}/\d{2}/\d{4})\s+(.+?)\s+(.+?)\s+(\d+)\s+([\d,]+\.\d{2})\s+(CR|DR)"
    )
    transactions = []
    categories = load_category_mappings()

    with pdfplumber.open(pdf_path, password=password) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            for line in text.splitlines():
                match = pattern.match(line.strip())
                if match:
                    date_str, desc, merchant_cat, points, amount_str, cr_dr = match.groups()
                    date = pd.to_datetime(date_str, format="%d/%m/%Y", errors='coerce').date()
                    amount = float(amount_str.replace(",", ""))
                    debit = credit = 0
                    if cr_dr.upper() == "DR":
                        debit = amount
                    else:
                        credit = amount

                    try:
                        points = int(points)
                    except Exception:
                        points = 0

                    category = assign_category(desc, merchant_cat, categories)

                    transactions.append([
                        date,
                        desc.strip(),
                        merchant_cat.strip(),
                        category.strip(),
                        points,
                        "IndusInd",
                        "CreditCard",
                        debit,
                        credit,
                        None
                    ])

    return pd.DataFrame(transactions, columns=[
        'Date', 'Description', 'Merchant', 'Category',
        'Reward Points', 'Bank', 'SourceType', 'Debit', 'Credit', 'Balance'
    ])


# cc_analysis/extractors.py

from openpyxl import load_workbook
import pandas as pd
import os
from cc_analysis.utils import log_info, log_error

DEFAULT_SHEET_NAME = "All_Transactions"
APPEND_KEY_COLUMNS = ['Date', 'Description', 'Debit', 'Credit', 'Bank', 'SourceType']

def append_to_excel(df, source_type, excel_path, sheet_name=DEFAULT_SHEET_NAME):
    try:
        final_cols = [
            'Date', 'Description', 'Merchant', 'Category',
            'Debit', 'Credit', 'Balance', 'Amount',
            'Reward Points', 'Bank', 'SourceType'
        ]

        if df.empty:
            log_info("⚠️ No data to append.")
            return

        # Reorder and fill missing columns
        df = df.copy()
        for col in final_cols:
            if col not in df.columns:
                df[col] = None
        df = df[final_cols]

        if not os.path.exists(excel_path):
            df.to_excel(excel_path, sheet_name=sheet_name, index=False, engine='openpyxl')
            log_info(f"📘 Excel file not found. Creating new file: {os.path.basename(excel_path)}")
        else:
            # Load existing data
            if sheet_name in load_workbook(excel_path).sheetnames:
                existing_df = pd.read_excel(excel_path, sheet_name=sheet_name, engine='openpyxl')
            else:
                existing_df = pd.DataFrame(columns=final_cols)

            # Combine and de-duplicate
            combined_df = pd.concat([existing_df, df], ignore_index=True)
            combined_df.drop_duplicates(subset=APPEND_KEY_COLUMNS, inplace=True)

            # Save back
            with pd.ExcelWriter(excel_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                combined_df.to_excel(writer, sheet_name=sheet_name, index=False)

            log_info(f"✅ Appended data to Excel: {excel_path} ({source_type}) | Final rows: {len(combined_df)}")

    except Exception as e:
        log_error(f"append_to_excel failed: {e}")


def save_to_consolidated_db(df, source_type, db_path=CONSOLIDATED_DB_PATH):
    if df.empty:
        log_error("save_to_consolidated_db: Empty dataframe.")
        return 0
    try:
        df = df.copy()
        df["SourceType"] = source_type
        if 'Amount' in df.columns:
            df.drop(columns=['Amount'], inplace=True)

        with sqlite3.connect(db_path) as conn:
            if conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='transactions'").fetchone():
                existing_df = pd.read_sql("SELECT * FROM transactions", conn)
            else:
                existing_df = pd.DataFrame(columns=df.columns)

            combined_df = pd.concat([existing_df, df], ignore_index=True)
            combined_df.drop_duplicates(subset=["Date", "Description", "Bank", "SourceType", "Debit", "Credit"], inplace=True)
            combined_df.to_sql("transactions", conn, if_exists="replace", index=False)
            return len(combined_df)
    except Exception as e:
        log_error(f"save_to_consolidated_db failed: {e}")
        return 0